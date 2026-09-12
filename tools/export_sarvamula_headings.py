#!/usr/bin/env python3
"""Every heading in DvaitaVedanta/SarvaMula, as a downloadable workbook.

WHAT A "COMMENTARY TITLE" IS HERE. The reader's .commentary-title element is
fed from two different places (js/render.js): the commentary's display name
for a layered grantha, and an analysis field's label for a per-unit block. Under
SarvaMula each data.json IS one layer, so the title that matters is the
document's own, and the headings a reader actually sees beneath it are the
`reference` hierarchy -- a ">"-joined path like

    कृष्णामृतमहार्णवः > मङ्गलाचरणम् > प्रथमोऽध्यायः

So the workbook carries both: one sheet of works, and one row per distinct
heading path with its levels split into columns, which is what makes it
sortable and filterable in Excel rather than a wall of strings.
"""

from __future__ import annotations

import argparse
import collections
import glob
import json
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
ROOT = REPO / "dge" / "data" / "DvaitaVedanta" / "SarvaMula"


def collect(root: Path):
    works, headings = [], []
    for f in sorted(root.glob("**/data.json")):
        rel = f.relative_to(root).parent.as_posix()
        try:
            d = json.loads(f.read_text(encoding="utf-8"))
        except Exception as e:
            works.append({"path": rel, "title": f"UNREADABLE: {e}", "author": "", "units": 0})
            continue
        units = d.get("items") or list((d.get("shlokas") or {}).values())
        works.append({
            "path": rel,
            "title": d.get("title") or "",
            "author": d.get("default_author") or "",
            "units": len(units),
        })
        # Distinct heading paths, in the order they first appear -- a reader
        # moving through the text meets them in this order, and sorting
        # alphabetically would lose that.
        seen = collections.OrderedDict()
        for u in units:
            if not isinstance(u, dict):
                continue
            ref = (u.get("reference") or "").strip()
            if not ref:
                continue
            key = ref
            if key not in seen:
                seen[key] = {"count": 0, "section": u.get("section") or "",
                             "unit_title": u.get("unit_title") or ""}
            seen[key]["count"] += 1
        for ref, info in seen.items():
            levels = [p.strip() for p in ref.split(">")]
            headings.append({
                "path": rel,
                "work": d.get("title") or "",
                "reference": ref,
                "level_1": levels[0] if len(levels) > 0 else "",
                "level_2": levels[1] if len(levels) > 1 else "",
                "level_3": levels[2] if len(levels) > 2 else "",
                "level_4": " > ".join(levels[3:]) if len(levels) > 3 else "",
                "section": info["section"],
                "unit_title": info["unit_title"],
                "units": info["count"],
            })
    return works, headings


def write_xlsx(works, headings, out: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    # A Devanagari-capable font: the default Calibri has no Devanagari and
    # every heading in this workbook would render as boxes.
    DEVA = "Nirmala UI"

    def sheet(ws, title, cols, rows, widths):
        ws.title = title
        ws.append([c[1] for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append([r.get(c[0], "") for c in cols])
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.font = Font(name=DEVA)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    sheet(wb.active, "Works",
          [("title", "Title"), ("author", "Author"), ("units", "Units"), ("path", "Folder")],
          works, [34, 34, 8, 52])

    sheet(wb.create_sheet(), "Headings",
          [("work", "Work"), ("level_1", "Level 1"), ("level_2", "Level 2"),
           ("level_3", "Level 3"), ("level_4", "Level 4+"), ("section", "Section"),
           ("unit_title", "Unit title"), ("units", "Units"), ("reference", "Full reference"),
           ("path", "Folder")],
          headings, [28, 26, 26, 26, 26, 24, 28, 8, 56, 48])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--root", default=str(ROOT))
    ap.add_argument("--out", default=str(REPO / "dge" / "data" / "catalogs" / "sarvamula_headings.xlsx"))
    args = ap.parse_args(argv)

    root = Path(args.root)
    if not root.is_dir():
        raise SystemExit(f"no corpus at {root}")
    works, headings = collect(root)
    write_xlsx(works, headings, Path(args.out))
    print(f"{len(works)} works, {len(headings)} distinct headings -> {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
